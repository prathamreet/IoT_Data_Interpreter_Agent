"""Generate the IoT Data Interpreter Agent pitch deck (python-pptx).

Run locally:   python build_deck.py [output.pptx]
Run in Docker: see deck/README.md  (recommended — no local Python needed)

Design notes:
- Committed dark, premium theme (navy + cyan signal accent) that mirrors the
  product's own dashboard — a content-informed palette, not generic blue.
- Engineering "comment" kicker motif (// NN - SECTION) repeated on every slide.
- Honest by construction: factual stat callouts only; no invented metrics; a
  dedicated "what's built vs. what's next" slide; evaluation stated as a method
  you reproduce, not numbers we claim.
"""

from __future__ import annotations

import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ── palette ───────────────────────────────────────────────────────────
NAVY = RGBColor(0x0A, 0x0F, 0x1F)
PANEL = RGBColor(0x12, 0x1A, 0x30)
PANEL2 = RGBColor(0x0F, 0x17, 0x29)
INK = RGBColor(0xE9, 0xEF, 0xF8)
MUTE = RGBColor(0x9A, 0xA7, 0xC0)
FAINT = RGBColor(0x60, 0x6F, 0x8A)
CYAN = RGBColor(0x38, 0xBD, 0xF8)
INDIGO = RGBColor(0x9D, 0xA6, 0xFB)
GREEN = RGBColor(0x34, 0xD3, 0x99)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
RED = RGBColor(0xF2, 0x6D, 0x6D)
LINE = RGBColor(0x25, 0x33, 0x4D)

HEAD = "Trebuchet MS"   # header font with personality
BODY = "Calibri"        # clean body font
MONO = "Consolas"       # engineering / telemetry accents

W, H = 13.333, 7.5      # 16:9 in inches
MARGIN = 0.75

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)


# ── primitives ────────────────────────────────────────────────────────
def slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid()
    r.fill.fore_color.rgb = NAVY
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def rect(s, x, y, w, h, fill, line=None, rounded=False, radius=0.07, line_w=1.0):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp


def text(s, x, y, w, h, lines, size=15, color=INK, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=BODY, gap=5,
         leading=1.06, margin=0.04):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(margin))
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        p.space_before = Pt(0)
        p.line_spacing = leading
        txt, o = (ln if isinstance(ln, tuple) else (ln, {}))
        run = p.add_run()
        run.text = txt
        f = run.font
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.italic = o.get("italic", italic)
        f.name = o.get("font", font)
        f.color.rgb = o.get("color", color)
    return tb


def bullets(s, x, y, w, h, items, size=14.5, gap=9, marker_color=CYAN,
            head_color=INK, body_color=MUTE, leading=1.08):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.04))
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = leading
        mk = p.add_run()
        mk.text = "—  "
        mk.font.name = MONO
        mk.font.size = Pt(size)
        mk.font.bold = True
        mk.font.color.rgb = marker_color
        if isinstance(it, tuple):
            head, body = it
            r1 = p.add_run()
            r1.text = head + " "
            r1.font.name = BODY
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = head_color
            r2 = p.add_run()
            r2.text = body
            r2.font.name = BODY
            r2.font.size = Pt(size)
            r2.font.color.rgb = body_color
        else:
            r = p.add_run()
            r.text = it
            r.font.name = BODY
            r.font.size = Pt(size)
            r.font.color.rgb = head_color
    return tb


def card(s, x, y, w, h, head, body, accent=None):
    rect(s, x, y, w, h, PANEL, line=LINE, rounded=True, radius=0.08)
    if accent:
        bar = rect(s, x, y + 0.18, 0.07, h - 0.36, accent, rounded=False)
        bar.line.fill.background()
    pad = 0.45 if accent else 0.26
    text(s, x + pad, y + 0.18, w - pad - 0.22, h - 0.34,
         [(head, {"size": 15, "bold": True, "color": INK, "font": HEAD}),
          (body, {"size": 12.5, "color": MUTE, "font": BODY})],
         anchor=MSO_ANCHOR.MIDDLE, gap=4)


def numdot(s, x, y, d, label, color=CYAN):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    o.fill.solid()
    o.fill.fore_color.rgb = PANEL2
    o.line.color.rgb = color
    o.line.width = Pt(1.5)
    o.shadow.inherit = False
    tf = o.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = MONO
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = color
    return o


def sparkline(s, x, y, w, h, pts, color=CYAN, width=2.25):
    n = len(pts)
    for i in range(n - 1):
        x1 = x + w * i / (n - 1)
        x2 = x + w * (i + 1) / (n - 1)
        y1 = y + h * (1 - pts[i])
        y2 = y + h * (1 - pts[i + 1])
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(width)
        c.shadow.inherit = False


def vconnect(s, x, y1, y2, color=FAINT, width=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y1), Inches(x), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    c.shadow.inherit = False


def chip(s, x, y, w, h, label, fg=INK, border=LINE, fill=PANEL, font=BODY, size=12, bold=False):
    rect(s, x, y, w, h, fill, line=border, rounded=True, radius=0.18)
    text(s, x + 0.12, y, w - 0.24, h, [(label, {"size": size, "color": fg, "font": font, "bold": bold})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def base(s, idx, kicker, title, title_size=31):
    text(s, MARGIN, 0.52, 11.9, 0.4,
         [(kicker, {"size": 12.5, "color": CYAN, "bold": True, "font": MONO})])
    text(s, MARGIN, 0.92, 11.9, 1.0,
         [(title, {"size": title_size, "color": INK, "bold": True, "font": HEAD})])
    text(s, MARGIN, H - 0.46, 8, 0.32,
         [("IoT Data Interpreter Agent", {"size": 9.5, "color": FAINT, "font": BODY})])
    text(s, W - 1.5, H - 0.46, 0.75, 0.32,
         [(f"{idx:02d}", {"size": 9.5, "color": FAINT, "font": MONO})], align=PP_ALIGN.RIGHT)


# ── slides ────────────────────────────────────────────────────────────
def s_title():
    s = slide()
    text(s, MARGIN, 1.35, 11, 0.4,
         [("INDUSTRIAL IoT  //  AI REASONING LAYER", {"size": 13, "color": CYAN, "bold": True, "font": MONO})])
    text(s, MARGIN, 1.85, 11.8, 1.4,
         [("IoT Data Interpreter Agent", {"size": 47, "color": INK, "bold": True, "font": HEAD})])
    text(s, MARGIN, 3.25, 9.7, 1.0,
         [("Turns raw sensor streams into explanations, anomaly flags, and "
           "maintenance actions — in real time.", {"size": 18.5, "color": MUTE, "font": BODY})],
         leading=1.18)
    text(s, MARGIN, 4.5, 10.5, 0.6,
         [("“The future of IoT is not more dashboards. "
           "It's machines that explain themselves.”",
           {"size": 15.5, "color": CYAN, "italic": True, "font": BODY})])
    sparkline(s, MARGIN, 5.35, 7.6, 0.95,
              [0.30, 0.34, 0.28, 0.33, 0.31, 0.36, 0.30, 0.92, 0.62, 0.55, 0.58], color=CYAN, width=2.5)
    chip(s, MARGIN, 6.62, 3.5, 0.5, "Builathon  ·  Use Case #43", fg=INK, font=BODY, size=12.5)
    chip(s, MARGIN + 3.7, 6.62, 4.25, 0.5, "Working prototype  ·  no hardware required",
         fg=GREEN, border=RGBColor(0x1E, 0x4D, 0x3E), font=BODY, size=12.5)
    text(s, W - 3.4, 6.66, 3.0, 0.4,
         [("Team / Presenter:  ____", {"size": 11.5, "color": FAINT, "font": MONO})], align=PP_ALIGN.RIGHT)


def s_problem():
    s = slide()
    base(s, 1, "// 01 - THE PROBLEM", "Dashboards show data, not meaning")
    text(s, MARGIN, 2.15, 5.55, 2.6,
         [("A red value tells an operator a threshold was crossed. It does not say "
           "why it happened, whether it matters, or what to do next.",
           {"size": 16.5, "color": INK, "font": BODY}),
          ("Interpreting every spike still needs a domain expert — and there are "
           "never enough to watch every sensor.",
           {"size": 16.5, "color": MUTE, "font": BODY})],
         leading=1.22, gap=12)
    rows = [
        ("Undetected anomalies", "Unplanned downtime and safety incidents", RED),
        ("Alert fatigue", "Real warnings ignored, response delayed", AMBER),
        ("Reactive maintenance", "Higher repair cost, shorter asset life", AMBER),
        ("Noisy or missing data", "Decisions made on corrupted signals", CYAN),
    ]
    cx, cw, ch, gap = 6.85, 5.75, 0.96, 0.18
    for i, (h, b, a) in enumerate(rows):
        card(s, cx, 2.15 + i * (ch + gap), cw, ch, h, b, accent=a)


def s_whynow():
    s = slide()
    base(s, 2, "// 02 - WHY NOW", "Three shifts make this possible")
    cols = [
        ("01", "LLM reasoning", "Models read structured sensor data and explain it in clear "
         "language — well beyond fixed alert templates."),
        ("02", "Mature IoT", "MQTT and managed IoT hubs make ingesting real sensor streams "
         "standard, cheap, and reliable."),
        ("03", "Agent tooling", "Tool-use loops let a model investigate across many signals, "
         "not just answer a single prompt."),
    ]
    cw, gap = 3.78, 0.28
    x0 = MARGIN
    for i, (n, h, b) in enumerate(cols):
        x = x0 + i * (cw + gap)
        rect(s, x, 2.35, cw, 3.5, PANEL, line=LINE, rounded=True, radius=0.06)
        numdot(s, x + 0.35, 2.75, 0.66, n, color=CYAN)
        text(s, x + 0.35, 3.65, cw - 0.7, 0.5,
             [(h, {"size": 19, "color": INK, "bold": True, "font": HEAD})])
        text(s, x + 0.35, 4.25, cw - 0.7, 1.4,
             [(b, {"size": 14, "color": MUTE, "font": BODY})], leading=1.22)


def s_solution():
    s = slide()
    base(s, 3, "// 03 - THE SOLUTION", "A reasoning layer between sensors and people")
    text(s, MARGIN, 1.95, 11.9, 0.4,
         [("Not a smarter chart. A layer that detects, explains, and recommends.",
           {"size": 15.5, "color": MUTE, "font": BODY, "italic": True})])
    labels = ["Sensor stream", "Clean + score quality", "Detect anomalies",
              "Narrate (LLM)", "Recommend + report"]
    cw, gap, ch, y = 2.0, 0.42, 1.0, 2.8
    x = MARGIN
    for i, lab in enumerate(labels):
        accent = CYAN if i in (2, 3) else PANEL
        rect(s, x, y, cw, ch, PANEL, line=(CYAN if i in (2, 3) else LINE), rounded=True, radius=0.12)
        text(s, x + 0.12, y, cw - 0.24, ch,
             [(lab, {"size": 13, "color": INK, "bold": (i in (2, 3)), "font": BODY})],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, leading=1.05)
        if i < len(labels) - 1:
            text(s, x + cw, y, gap, ch, [(">", {"size": 20, "color": CYAN, "bold": True, "font": MONO})],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += cw + gap
    # Honest, factual numbers (no invented metrics).
    stats = [("6", "sensors simulated"), ("3", "detection methods"),
             ("5", "fault scenarios"), ("0", "API keys to run")]
    sw = 11.9 / 4
    for i, (n, lab) in enumerate(stats):
        sx = MARGIN + i * sw
        text(s, sx, 4.7, sw - 0.2, 0.8, [(n, {"size": 46, "color": CYAN, "bold": True, "font": HEAD})])
        text(s, sx, 5.65, sw - 0.2, 0.5, [(lab, {"size": 13, "color": MUTE, "font": BODY})])


def s_built():
    s = slide()
    base(s, 4, "// 04 - WHAT'S BUILT", "A working prototype, end to end")
    left = [
        "Realistic multi-sensor simulator with switchable fault scenarios",
        "Multi-method anomaly detection (statistical + ML)",
        "Data-quality scoring: broken sensor vs. real fault",
        "LLM narration with a deterministic fallback",
    ]
    right = [
        "Agentic investigator: a real Claude tool-use loop",
        "Live dashboard — no front-end build, no CDN",
        "Auto-generated incident report (HTML / PDF)",
        "Reproducible evaluation harness; fully Dockerized",
    ]
    bullets(s, MARGIN, 2.15, 5.7, 2.7, left, size=14.5, gap=11)
    bullets(s, MARGIN + 6.05, 2.15, 5.7, 2.7, right, size=14.5, gap=11)
    rect(s, MARGIN, 5.55, 11.85, 1.05, PANEL2, line=LINE, rounded=True, radius=0.06)
    text(s, MARGIN + 0.3, 5.6, 11.3, 0.95,
         [("Honest scope:", {"size": 14, "color": AMBER, "bold": True, "font": HEAD}),
          ("There is no physical hardware. Sensors are simulated. Real devices integrate "
           "behind a single interface — see the hardware slide.",
           {"size": 13.5, "color": MUTE, "font": BODY})],
         anchor=MSO_ANCHOR.MIDDLE, gap=3)


def s_demo():
    s = slide()
    base(s, 5, "// 05 - LIVE DEMO", "Sixty seconds, four moments")
    steps = [
        ("01", "Trigger an HVAC failure", "Temperature climbs; a narrated anomaly appears, "
         "with a correlated humidity sensor."),
        ("02", "Trigger a sensor malfunction", "The system flags a sensor fault — not a "
         "false equipment alarm."),
        ("03", "Run Deep Investigation", "The agent inspects the fleet and writes a "
         "prioritized incident report."),
        ("04", "Export the report", "One click to a clean HTML report, ready to save as PDF."),
    ]
    cw, ch, gx, gy = 5.78, 1.62, 0.28, 0.24
    for i, (n, h, b) in enumerate(steps):
        x = MARGIN + (i % 2) * (cw + gx)
        y = 2.2 + (i // 2) * (ch + gy)
        rect(s, x, y, cw, ch, PANEL, line=LINE, rounded=True, radius=0.06)
        numdot(s, x + 0.3, y + 0.3, 0.6, n, color=CYAN)
        text(s, x + 1.1, y + 0.26, cw - 1.35, 0.5,
             [(h, {"size": 16.5, "color": INK, "bold": True, "font": HEAD})])
        text(s, x + 1.1, y + 0.74, cw - 1.35, 0.8,
             [(b, {"size": 13, "color": MUTE, "font": BODY})], leading=1.16)
    text(s, MARGIN, 6.35, 11.9, 0.4,
         [("Runs fully offline. With an API key, the same flow uses live Claude.",
           {"size": 13.5, "color": CYAN, "font": BODY, "italic": True})])


def s_architecture():
    s = slide()
    base(s, 6, "// 06 - ARCHITECTURE", "How it works")
    boxes = [
        ("Data source", "simulator  |  MQTT (real hardware)", "sources/"),
        ("Processing", "buffer  ->  impute  ->  quality  ->  anomaly", "processing/"),
        ("Engine", "orchestration  ·  edge-triggered alerts  ·  live WebSocket", "engine.py"),
        ("Reasoning", "narrator (1 grounded call)  +  investigator (agent loop)  +  fallback", "agent/"),
        ("Delivery", "live dashboard  +  incident report", "frontend · reporting"),
    ]
    bw, bh, gap, y = 8.5, 0.74, 0.28, 2.05
    bx = (W - bw) / 2 - 0.9
    for i, (h, b, mod) in enumerate(boxes):
        accent = CYAN if i in (1, 3) else INDIGO
        rect(s, bx, y, bw, bh, PANEL, line=LINE, rounded=True, radius=0.09)
        rb = rect(s, bx, y + 0.14, 0.07, bh - 0.28, accent)
        rb.line.fill.background()
        text(s, bx + 0.35, y, 2.5, bh,
             [(h, {"size": 15, "color": INK, "bold": True, "font": HEAD})],
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, bx + 2.7, y, bw - 2.9, bh,
             [(b, {"size": 12.5, "color": MUTE, "font": MONO})],
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, bx + bw + 0.25, y, 2.4, bh,
             [(mod, {"size": 12, "color": CYAN, "font": MONO})], anchor=MSO_ANCHOR.MIDDLE)
        if i < len(boxes) - 1:
            vconnect(s, bx + bw / 2, y + bh, y + bh + gap, color=FAINT, width=1.5)
        y += bh + gap


def s_detection():
    s = slide()
    base(s, 7, "// 07 - DETECTION", "Detection you can defend")
    items = [
        ("Z-score vs. a reference window —", "catches gradual ramps, not just spikes."),
        ("IQR fence —", "robust outlier bounds, resistant to skew."),
        ("Isolation Forest —", "unsupervised ML over value and local slope."),
        ("Voting + 2-tick persistence —", "sensitive to real faults, quiet on noise."),
        ("Confidence scoring —", "flags low-quality data instead of trusting it."),
    ]
    bullets(s, MARGIN, 2.2, 7.0, 4.0, items, size=15.5, gap=15)
    # Visual: a signal that ramps then spikes, plus a severity legend.
    px, pw = 8.35, 4.25
    rect(s, px, 2.2, pw, 2.5, PANEL, line=LINE, rounded=True, radius=0.06)
    text(s, px + 0.3, 2.35, pw - 0.6, 0.4,
         [("LIVE SIGNAL", {"size": 11.5, "color": FAINT, "bold": True, "font": MONO})])
    sparkline(s, px + 0.35, 2.95, pw - 0.7, 1.45,
              [0.25, 0.27, 0.24, 0.30, 0.45, 0.6, 0.78, 0.9, 0.84, 0.88], color=CYAN, width=2.5)
    legend = [("Low", GREEN), ("Moderate", AMBER), ("Critical", RED)]
    lx = px
    for lab, col in legend:
        sq = rect(s, lx, 5.05, 0.20, 0.20, col, rounded=False)
        sq.line.fill.background()
        text(s, lx + 0.30, 4.98, 1.0, 0.35, [(lab, {"size": 11.5, "color": MUTE, "font": BODY})])
        lx += 1.3


def s_agentic():
    s = slide()
    base(s, 8, "// 08 - THE AGENTIC DIFFERENCE", "From narration to investigation")
    panels = [
        ("Narrator", "One grounded, structured call per anomaly:",
         ["plain-English summary", "severity (Low / Moderate / Critical)",
          "root-cause hypothesis", "recommended action"], CYAN),
        ("Investigator", "A real tool-use loop across the fleet:",
         ["list sensors", "pull recent history", "find correlations",
          "write a prioritized incident report"], INDIGO),
    ]
    pw, gap = 5.78, 0.29
    for i, (h, sub, pts, col) in enumerate(panels):
        x = MARGIN + i * (pw + gap)
        rect(s, x, 2.15, pw, 3.55, PANEL, line=LINE, rounded=True, radius=0.05)
        text(s, x + 0.4, 2.4, pw - 0.8, 0.5, [(h, {"size": 21, "color": col, "bold": True, "font": HEAD})])
        text(s, x + 0.4, 2.98, pw - 0.8, 0.5, [(sub, {"size": 13.5, "color": INK, "font": BODY})])
        bullets(s, x + 0.4, 3.5, pw - 0.8, 2.0, pts, size=14, gap=8, marker_color=col)
    rect(s, MARGIN, 5.95, 11.85, 0.7, PANEL2, line=LINE, rounded=True, radius=0.08)
    text(s, MARGIN + 0.3, 5.95, 11.3, 0.7,
         [("Grounded strictly in the data. The system says “undetermined” rather than guessing.",
           {"size": 13.5, "color": MUTE, "font": BODY, "italic": True})], anchor=MSO_ANCHOR.MIDDLE)


def s_trust():
    s = slide()
    base(s, 9, "// 09 - TRUST", "Built to be trusted")
    rows = [
        ("Sensor fault vs. real anomaly", "Confidence scoring separates a malfunctioning sensor "
         "from a genuine process change."),
        ("Communicates uncertainty", "Every sensor carries a confidence score; low-quality data "
         "is flagged, not trusted."),
        ("Grounded by design", "Prompts forbid speculation beyond the data, preventing "
         "hallucinated causes."),
        ("Works without the cloud", "A deterministic fallback keeps the product running with no "
         "API key and no network."),
    ]
    y = 2.2
    for i, (h, b) in enumerate(rows):
        rect(s, MARGIN, y, 11.85, 0.96, PANEL, line=LINE, rounded=True, radius=0.07)
        numdot(s, MARGIN + 0.3, y + 0.2, 0.56, f"{i+1:02d}", color=CYAN)
        text(s, MARGIN + 1.15, y, 3.6, 0.96,
             [(h, {"size": 15.5, "color": INK, "bold": True, "font": HEAD})], anchor=MSO_ANCHOR.MIDDLE)
        text(s, MARGIN + 4.95, y, 6.7, 0.96,
             [(b, {"size": 13.5, "color": MUTE, "font": BODY})], anchor=MSO_ANCHOR.MIDDLE, leading=1.12)
        y += 0.96 + 0.16


def s_hardware():
    s = slide()
    base(s, 10, "// 10 - HARDWARE-READY", "Hardware is a drop-in, not a rebuild")
    text(s, MARGIN, 2.2, 6.0, 2.0,
         [("Everything consumes one DataSource interface. The simulator is one "
           "implementation; MQTT is another.", {"size": 16.5, "color": INK, "font": BODY})],
         leading=1.25)
    steps = ["Point sensors at an MQTT broker", "Set one variable: DATA_SOURCE=mqtt",
             "Adapt one parser to your schema"]
    y = 3.7
    for i, st in enumerate(steps):
        numdot(s, MARGIN, y, 0.5, f"{i+1}", color=CYAN)
        text(s, MARGIN + 0.7, y - 0.05, 5.3, 0.6, [(st, {"size": 14.5, "color": MUTE, "font": BODY})],
             anchor=MSO_ANCHOR.MIDDLE)
        y += 0.72
    # code chip panel
    px, pw = 7.1, 5.45
    rect(s, px, 2.2, pw, 2.35, PANEL2, line=LINE, rounded=True, radius=0.05)
    text(s, px + 0.35, 2.4, pw - 0.7, 0.4,
         [(".env", {"size": 12, "color": FAINT, "bold": True, "font": MONO})])
    text(s, px + 0.35, 2.95, pw - 0.7, 1.4,
         [("DATA_SOURCE=mqtt", {"size": 16, "color": CYAN, "font": MONO}),
          ("MQTT_TOPIC=sensors/#", {"size": 16, "color": INK, "font": MONO})], gap=8)
    text(s, px, 4.75, pw, 1.0,
         [("No other code changes.", {"size": 15, "color": GREEN, "bold": True, "font": HEAD}),
          ("Same pipeline — real telemetry instead of simulated.", {"size": 13.5, "color": MUTE, "font": BODY})],
         gap=4)


def s_eval():
    s = slide()
    base(s, 11, "// 11 - EVALUATION", "Measured, not claimed")
    text(s, MARGIN, 2.2, 11.6, 2.0,
         [("A labeled golden dataset and an evaluation harness report precision, recall, "
           "and F1 for the detector — reproducible with a single command.",
           {"size": 17, "color": INK, "font": BODY}),
          ("The dataset deliberately includes a stuck-sensor case, to confirm the system "
           "does not raise a process anomaly for a malfunction.",
           {"size": 16, "color": MUTE, "font": BODY})], leading=1.24, gap=12)
    rect(s, MARGIN, 4.35, 6.2, 0.8, PANEL2, line=LINE, rounded=True, radius=0.1)
    text(s, MARGIN + 0.3, 4.35, 5.8, 0.8,
         [("python -m evaluation.evaluate", {"size": 16, "color": CYAN, "font": MONO})],
         anchor=MSO_ANCHOR.MIDDLE)
    rect(s, MARGIN, 5.5, 11.85, 0.95, PANEL, line=LINE, rounded=True, radius=0.07)
    text(s, MARGIN + 0.3, 5.5, 11.3, 0.95,
         [("We present the method and let the numbers come from your machine — "
           "not metrics we invented for a slide.",
           {"size": 14, "color": MUTE, "font": BODY, "italic": True})], anchor=MSO_ANCHOR.MIDDLE)


def s_roadmap():
    s = slide()
    base(s, 12, "// 12 - ROADMAP", "What's built, what's next")
    built = ["Sensor simulation + scenarios", "Multi-method detection", "Narration + fallback",
             "Agentic investigator", "Dashboard + report", "Evaluation harness", "Docker + hardware interface"]
    nxt = ["Validate against real hardware streams", "Historical store + trend baselining",
           "Human feedback to tune recommendations", "Alert routing (Slack / SMS / email)",
           "Auth + multi-tenant", "Scale-out ingestion"]
    pw, gap = 5.78, 0.29
    rect(s, MARGIN, 2.15, pw, 4.2, PANEL, line=LINE, rounded=True, radius=0.05)
    text(s, MARGIN + 0.4, 2.38, pw - 0.8, 0.5, [("BUILT", {"size": 15, "color": GREEN, "bold": True, "font": MONO})])
    bullets(s, MARGIN + 0.4, 2.95, pw - 0.8, 3.3, built, size=14, gap=8, marker_color=GREEN)
    x2 = MARGIN + pw + gap
    rect(s, x2, 2.15, pw, 4.2, PANEL, line=LINE, rounded=True, radius=0.05)
    text(s, x2 + 0.4, 2.38, pw - 0.8, 0.5, [("NEXT", {"size": 15, "color": INDIGO, "bold": True, "font": MONO})])
    bullets(s, x2 + 0.4, 2.95, pw - 0.8, 3.3, nxt, size=14, gap=8, marker_color=INDIGO)


def s_close():
    s = slide()
    text(s, MARGIN, 1.5, 11, 0.4,
         [("// 13 - THE ASK", {"size": 13, "color": CYAN, "bold": True, "font": MONO})])
    text(s, MARGIN, 2.0, 11.8, 1.2,
         [("Machines that explain themselves", {"size": 40, "color": INK, "bold": True, "font": HEAD})])
    text(s, MARGIN, 3.3, 10.8, 1.0,
         [("A reasoning layer that makes industrial intelligence accessible to every "
           "operator — not just data scientists.", {"size": 17.5, "color": MUTE, "font": BODY})],
         leading=1.22)
    rect(s, MARGIN, 4.55, 11.85, 1.5, PANEL, line=LINE, rounded=True, radius=0.05)
    text(s, MARGIN + 0.4, 4.7, 4.0, 1.2,
         [("What we're looking for", {"size": 15, "color": CYAN, "bold": True, "font": HEAD})],
         anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, MARGIN + 4.3, 4.72, 7.2, 1.2,
            ["A pilot sensor feed or dataset to validate against",
             "Mentorship to harden it for production"], size=14.5, gap=8)
    sparkline(s, MARGIN, 6.35, 7.6, 0.5,
              [0.3, 0.34, 0.3, 0.36, 0.32, 0.9, 0.6, 0.58], color=CYAN, width=2.25)
    text(s, W - 4.0, 6.4, 3.6, 0.4,
         [("Team ____  ·  contact ____", {"size": 12, "color": FAINT, "font": MONO})], align=PP_ALIGN.RIGHT)


def build(path: str):
    s_title()
    s_problem()
    s_whynow()
    s_solution()
    s_built()
    s_demo()
    s_architecture()
    s_detection()
    s_agentic()
    s_trust()
    s_hardware()
    s_eval()
    s_roadmap()
    s_close()
    cp = prs.core_properties
    cp.title = "IoT Data Interpreter Agent"
    cp.author = "Builathon Use Case #43"
    cp.subject = "AI reasoning layer for IoT sensor streams"
    prs.save(path)
    print(f"Saved deck -> {path}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "IoT_Data_Interpreter_Pitch.pptx"
    build(out)
